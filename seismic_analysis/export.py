from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter


def create_excel_file(dataframe):
    """
    Create a formatted Excel workbook from the analysis results.
    """

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Analysis Results"

    # Write headers
    for column_index, column_name in enumerate(
        dataframe.columns,
        start=1,
    ):
        cell = worksheet.cell(
            row=1,
            column=column_index,
            value=column_name,
        )

        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # Write data
    for row_index, row in enumerate(
        dataframe.itertuples(index=False, name=None),
        start=2,
    ):
        for column_index, value in enumerate(
            row,
            start=1,
        ):
            cell = worksheet.cell(
                row=row_index,
                column=column_index,
                value=value,
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

    # Borders for all cells
    thin_side = Side(
        style="thin",
        color="000000",
    )

    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for row in worksheet.iter_rows():
        for cell in row:
            cell.border = border

    # Automatically adjust column widths
    for column_cells in worksheet.columns:
        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value)),
                )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 2, 12),
            30,
        )

    # Slightly taller header
    worksheet.row_dimensions[1].height = 25

    # Freeze header row
    worksheet.freeze_panes = "A2"

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output.getvalue()